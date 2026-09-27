"""
ingestion/loader.py
-------------------
Source Ingestion Module for Government of India / DGCI&S Road E-Way Bill workbook.

Governance Constraints:
1. Never modifies the source workbook (read-only mode).
2. Preserves source blanks/NULLs in raw extractions (no global 0 imputation).
3. Validates the five expected worksheets.
4. Records rich ingestion metadata per worksheet and run.
"""

import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
import openpyxl
import pandas as pd

from src.reconciliation.config import ReconciliationConfig, DEFAULT_CONFIG


class IngestionError(Exception):
    """Raised when source workbook cannot be ingested or validated."""
    pass


class MissingWorksheetError(IngestionError):
    """Raised when one or more expected worksheets are missing."""
    pass


class EwayBillSourceLoader:
    """Loads and extracts raw tables and metadata from the Road E-Way Bill workbook."""

    def __init__(self, file_path: str = "data/Road_EwayBill_2023_24.xlsx", config: Optional[ReconciliationConfig] = None):
        self.file_path = Path(file_path)
        self.config = config or DEFAULT_CONFIG
        self.ingestion_timestamp: Optional[datetime] = None
        self.metadata: Dict[str, Any] = {}

    def discover_and_validate(self) -> List[str]:
        """Verifies file existence and validates that all 5 expected worksheets exist."""
        if not self.file_path.exists():
            raise FileNotFoundError(f"Workbook not found at: {self.file_path.resolve()}")

        # Open in read-only mode to prevent any file modification
        wb = openpyxl.load_workbook(str(self.file_path), read_only=True, data_only=True)
        sheet_names = wb.sheetnames
        wb.close()

        missing = [s for s in self.config.expected_worksheets if s not in sheet_names]
        if missing:
            raise MissingWorksheetError(f"Missing expected worksheets: {missing}. Found: {sheet_names}")

        return sheet_names

    def load_raw_workbook(self) -> Dict[str, Any]:
        """
        Executes full source ingestion:
        - Validates sheets
        - Extracts raw tables with preserved NULLs
        - Compiles ingestion metadata
        """
        self.discover_and_validate()
        self.ingestion_timestamp = datetime.now(timezone.utc)

        # Load data workbook in data_only mode (read values, not formula strings)
        wb = openpyxl.load_workbook(str(self.file_path), data_only=True)

        raw_data = {
            "metadata": {
                "source_filename": self.file_path.name,
                "source_filepath": str(self.file_path.resolve()),
                "ingestion_timestamp": self.ingestion_timestamp.isoformat(),
                "worksheets": {}
            },
            "tables": {}
        }

        # 1. Sheet 1: Tab I_Stat_to_Stat_Revised_Road
        s1_data, s1_meta = self._extract_sheet_1(wb["Tab I_Stat_to_Stat_Revised_Road"])
        raw_data["tables"]["raw_state_to_state"] = s1_data
        raw_data["metadata"]["worksheets"]["Tab I_Stat_to_Stat_Revised_Road"] = s1_meta

        # 2. Sheet 2: Tab II_Chap_Revised_Road
        s2_data, s2_meta = self._extract_sheet_2(wb["Tab II_Chap_Revised_Road"])
        raw_data["tables"]["raw_chapter_summary"] = s2_data
        raw_data["metadata"]["worksheets"]["Tab II_Chap_Revised_Road"] = s2_meta

        # 3. Sheet 3: Tab III_Outward_Revised_Road
        s3_data, s3_meta = self._extract_sheet_3(wb["Tab III_Outward_Revised_Road"])
        raw_data["tables"]["raw_chapter_outward"] = s3_data
        raw_data["metadata"]["worksheets"]["Tab III_Outward_Revised_Road"] = s3_meta

        # 4. Sheet 4: Tab IV_Inward_Revised_Road
        s4_data, s4_meta = self._extract_sheet_4(wb["Tab IV_Inward_Revised_Road"])
        raw_data["tables"]["raw_chapter_inward"] = s4_data
        raw_data["metadata"]["worksheets"]["Tab IV_Inward_Revised_Road"] = s4_meta

        # 5. Sheet 5: Tab V_Internal_Revised_Road
        s5_data, s5_meta = self._extract_sheet_5(wb["Tab V_Internal_Revised_Road"])
        raw_data["tables"]["raw_chapter_internal"] = s5_data
        raw_data["metadata"]["worksheets"]["Tab V_Internal_Revised_Road"] = s5_meta

        wb.close()
        self.metadata = raw_data["metadata"]
        return raw_data

    def _extract_sheet_1(self, ws) -> Tuple[pd.DataFrame, Dict[str, Any]]:
        """Extracts Table I: State-to-State Matrix (Rows 4-36, Cols C-AI)."""
        # Column headers (From State) are in Row 2, Cols 3-35
        col_states = [str(ws.cell(2, c).value).strip() for c in range(3, 36)]
        
        records = []
        for r in range(4, 37):
            to_state = str(ws.cell(r, 2).value).strip()
            row_dict = {"to_state": to_state}
            for c_idx, from_state in enumerate(col_states, start=3):
                val = ws.cell(r, c_idx).value
                # Preserve raw NULL/None
                row_dict[from_state] = float(val) if val is not None else None
            records.append(row_dict)

        df = pd.DataFrame(records)
        meta = {
            "worksheet_name": "Tab I_Stat_to_Stat_Revised_Road",
            "title": ws.cell(1, 2).value,
            "data_start_row": 4,
            "data_end_row": 36,
            "data_start_col": 3,
            "data_end_col": 35,
            "row_count": len(df),
            "column_count": len(col_states) + 1,  # to_state + 33 states
            "null_cells_count": int(df.isna().sum().sum()),
            "states": col_states
        }
        return df, meta

    def _extract_sheet_2(self, ws) -> Tuple[pd.DataFrame, Dict[str, Any]]:
        """Extracts Table II: Chapter Summary (Rows 3-92, Cols B, C, D) and Row 93 Total."""
        records = []
        for r in range(3, 93):
            code = str(ws.cell(r, 2).value).strip()
            desc = str(ws.cell(r, 3).value).strip()
            val = ws.cell(r, 4).value
            records.append({
                "chapter_code": code,
                "chapter_description": desc,
                "value_inr_crore": float(val) if val is not None else None
            })

        df = pd.DataFrame(records)
        reported_total = ws.cell(93, 4).value
        meta = {
            "worksheet_name": "Tab II_Chap_Revised_Road",
            "title": ws.cell(1, 2).value,
            "data_start_row": 3,
            "data_end_row": 92,
            "row_count": len(df),
            "column_count": 3,
            "reported_total_row_93": float(reported_total) if reported_total is not None else None,
            "null_cells_count": int(df.isna().sum().sum())
        }
        return df, meta

    def _extract_sheet_3(self, ws) -> Tuple[pd.DataFrame, Dict[str, Any]]:
        """Extracts Table III: Chapter-wise Outward Movement (Rows 3-92, Cols B, C, and D-AJ)."""
        states = [str(ws.cell(2, c).value).strip() for c in range(4, 37)]
        records = []
        for r in range(3, 93):
            code = str(ws.cell(r, 2).value).strip()
            desc = str(ws.cell(r, 3).value).strip()
            row_dict = {
                "chapter_code": code,
                "chapter_description": desc
            }
            for c_idx, state_name in enumerate(states, start=4):
                val = ws.cell(r, c_idx).value
                row_dict[state_name] = float(val) if val is not None else None
            records.append(row_dict)

        df = pd.DataFrame(records)
        reported_totals = {
            state_name: float(ws.cell(93, c_idx).value) if ws.cell(93, c_idx).value is not None else None
            for c_idx, state_name in enumerate(states, start=4)
        }
        meta = {
            "worksheet_name": "Tab III_Outward_Revised_Road",
            "title": ws.cell(1, 2).value,
            "data_start_row": 3,
            "data_end_row": 92,
            "row_count": len(df),
            "column_count": len(states) + 2,
            "states": states,
            "reported_state_totals_row_93": reported_totals,
            "null_cells_count": int(df[states].isna().sum().sum())
        }
        return df, meta

    def _extract_sheet_4(self, ws) -> Tuple[pd.DataFrame, Dict[str, Any]]:
        """Extracts Table IV: Chapter-wise Inward Movement (Rows 3-92, Cols B, C, D-AJ, and AK)."""
        states = [str(ws.cell(2, c).value).strip() for c in range(4, 37)]
        records = []
        for r in range(3, 93):
            code = str(ws.cell(r, 2).value).strip()
            desc = str(ws.cell(r, 3).value).strip()
            row_dict = {
                "chapter_code": code,
                "chapter_description": desc
            }
            for c_idx, state_name in enumerate(states, start=4):
                val = ws.cell(r, c_idx).value
                row_dict[state_name] = float(val) if val is not None else None
            tot_val = ws.cell(r, 37).value
            row_dict["TOTAL"] = float(tot_val) if tot_val is not None else None
            records.append(row_dict)

        df = pd.DataFrame(records)
        reported_totals = {
            state_name: float(ws.cell(93, c_idx).value) if ws.cell(93, c_idx).value is not None else None
            for c_idx, state_name in enumerate(states, start=4)
        }
        grand_total = ws.cell(93, 37).value
        meta = {
            "worksheet_name": "Tab IV_Inward_Revised_Road",
            "title": ws.cell(1, 2).value,
            "data_start_row": 3,
            "data_end_row": 92,
            "row_count": len(df),
            "column_count": len(states) + 3,
            "states": states,
            "reported_state_totals_row_93": reported_totals,
            "reported_grand_total_cell_ak93": float(grand_total) if grand_total is not None else None,
            "null_cells_count": int(df[states].isna().sum().sum())
        }
        return df, meta

    def _extract_sheet_5(self, ws) -> Tuple[pd.DataFrame, Dict[str, Any]]:
        """Extracts Table V: Chapter-wise Internal Movement (Rows 3-92, Cols B, C, D-AJ, and AK)."""
        states = [str(ws.cell(2, c).value).strip() for c in range(4, 37)]
        records = []
        for r in range(3, 93):
            code = str(ws.cell(r, 2).value).strip()
            desc = str(ws.cell(r, 3).value).strip()
            row_dict = {
                "chapter_code": code,
                "chapter_description": desc
            }
            for c_idx, state_name in enumerate(states, start=4):
                val = ws.cell(r, c_idx).value
                row_dict[state_name] = float(val) if val is not None else None
            tot_val = ws.cell(r, 37).value
            row_dict["TOTAL"] = float(tot_val) if tot_val is not None else None
            records.append(row_dict)

        df = pd.DataFrame(records)
        reported_totals = {
            state_name: float(ws.cell(93, c_idx).value) if ws.cell(93, c_idx).value is not None else None
            for c_idx, state_name in enumerate(states, start=4)
        }
        grand_total = ws.cell(93, 37).value
        meta = {
            "worksheet_name": "Tab V_Internal_Revised_Road",
            "title": ws.cell(1, 2).value,
            "data_start_row": 3,
            "data_end_row": 92,
            "row_count": len(df),
            "column_count": len(states) + 3,
            "states": states,
            "reported_state_totals_row_93": reported_totals,
            "reported_grand_total_cell_ak93": float(grand_total) if grand_total is not None else None,
            "null_cells_count": int(df[states].isna().sum().sum())
        }
        return df, meta
