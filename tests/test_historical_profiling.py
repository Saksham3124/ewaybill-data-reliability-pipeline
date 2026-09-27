"""
tests/test_historical_profiling.py
----------------------------------
Unit tests for Historical Source Profiling of FY 2022-23 workbook
and structural comparison against FY 2023-24.

Verifies:
- 2022-23 workbook discovery & immutability
- five expected worksheets presence
- expected HS chapter coverage (90 chapters, 10 to 99)
- expected 33 jurisdiction extraction & label differences
- structural comparison (e.g. Table III Column AK availability)
- year identification in titles
"""

import hashlib
import os
from pathlib import Path
import openpyxl
import pytest

from src.reconciliation.config import ReconciliationConfig

F22_PATH = "data/Road_EwayBill_2022_23.xlsx"
F24_PATH = "data/Road_EwayBill_2023_24.xlsx"

EXPECTED_HASH_22 = "534ae64cdfe76ae1adbe5db789443cd5af5fc34df94925beba1949859d209aef"
EXPECTED_HASH_24 = "42fdba9a6fcf40fb47f9a632d403b28680e610160db51cf75511163f59ce803d"


def compute_sha256(filepath: str) -> str:
    h = hashlib.sha256()
    with open(filepath, "rb") as f:
        while chunk := f.read(8192):
            h.update(chunk)
    return h.hexdigest()


# -----------------------------------------------------------------------------
# 1. Workbook Discovery & Immutability
# -----------------------------------------------------------------------------

def test_2022_23_workbook_discovery():
    """Verifies that the historical 2022-23 workbook exists and is readable."""
    p = Path(F22_PATH)
    assert p.exists(), f"File not found: {F22_PATH}"
    assert p.is_file()
    assert p.stat().st_size == 173969


def test_source_immutability():
    """Verifies that both 2022-23 and 2023-24 workbooks remain unmodified (SHA-256 match)."""
    h22 = compute_sha256(F22_PATH)
    assert h22 == EXPECTED_HASH_22, f"2022-23 workbook hash mismatch! Expected {EXPECTED_HASH_22}, got {h22}"

    h24 = compute_sha256(F24_PATH)
    assert h24 == EXPECTED_HASH_24, f"2023-24 workbook hash mismatch! Expected {EXPECTED_HASH_24}, got {h24}"


# -----------------------------------------------------------------------------
# 2. Five Worksheets Presence
# -----------------------------------------------------------------------------

def test_five_worksheet_presence():
    """Verifies that all 5 expected worksheets exist in 2022-23 workbook."""
    wb = openpyxl.load_workbook(F22_PATH, read_only=True)
    expected_sheets = list(ReconciliationConfig().expected_worksheets)
    assert wb.sheetnames == expected_sheets
    wb.close()


# -----------------------------------------------------------------------------
# 3. Chapter Coverage & Descriptions
# -----------------------------------------------------------------------------

def test_expected_chapter_coverage():
    """Verifies that Table II in 2022-23 contains exactly 90 chapters (10-99) matching 2023-24 descriptions."""
    wb22 = openpyxl.load_workbook(F22_PATH, data_only=True)
    wb24 = openpyxl.load_workbook(F24_PATH, data_only=True)

    ws22 = wb22["Tab II_Chap_Revised_Road"]
    ws24 = wb24["Tab II_Chap_Revised_Road"]

    ch22 = [(str(ws22.cell(r, 2).value), str(ws22.cell(r, 3).value).strip()) for r in range(3, 93)]
    ch24 = [(str(ws24.cell(r, 2).value), str(ws24.cell(r, 3).value).strip()) for r in range(3, 93)]

    assert len(ch22) == 90
    assert ch22[0][0] == "10"
    assert ch22[-1][0] == "99"
    # Descriptions must be 100% identical
    assert ch22 == ch24

    wb22.close()
    wb24.close()


# -----------------------------------------------------------------------------
# 4. Jurisdiction Extraction & Label Differences
# -----------------------------------------------------------------------------

def test_expected_jurisdiction_extraction():
    """Verifies extraction of exactly 33 jurisdictions and confirms label differences."""
    wb22 = openpyxl.load_workbook(F22_PATH, data_only=True)
    ws1_22 = wb22["Tab I_Stat_to_Stat_Revised_Road"]

    # Origin states from row 2
    origins_22 = [ws1_22.cell(2, c).value for c in range(3, 36)]
    assert len(origins_22) == 33

    # In 2022-23: 'CHHATTISGARH' (two Hs) and 'JAMMU AND KASHMIR' (word AND)
    assert "CHHATTISGARH" in origins_22
    assert "CHATTISGARH" not in origins_22

    assert "JAMMU AND KASHMIR" in origins_22
    assert "JAMMU & KASHMIR" not in origins_22

    # In Table V of 2022-23, 'Other Territory' is Title Case
    ws5_22 = wb22["Tab V_Internal_Revised_Road"]
    col_26_val = ws5_22.cell(2, 26).value
    assert col_26_val == "Other Territory"

    wb22.close()


# -----------------------------------------------------------------------------
# 5. Structural Comparison (Table III Column AK)
# -----------------------------------------------------------------------------

def test_structural_comparison():
    """Verifies the key structural difference: Table III has Column AK in 2022-23, absent in 2023-24."""
    wb22 = openpyxl.load_workbook(F22_PATH, data_only=True)
    wb24 = openpyxl.load_workbook(F24_PATH, data_only=True)

    ws3_22 = wb22["Tab III_Outward_Revised_Road"]
    ws3_24 = wb24["Tab III_Outward_Revised_Road"]

    # 2022-23 has Column AK header in row 2
    assert ws3_22.cell(2, 37).value == "VALUE (in INR Crore)"
    assert ws3_22.cell(93, 37).value is not None
    assert float(ws3_22.cell(93, 37).value) > 60000000.0

    # 2023-24 has Column AK completely unpopulated
    assert ws3_24.cell(2, 37).value is None
    assert ws3_24.cell(93, 37).value is None

    wb22.close()
    wb24.close()


# -----------------------------------------------------------------------------
# 6. Year Identification in Titles
# -----------------------------------------------------------------------------

def test_year_identification():
    """Verifies that worksheet titles explicitly declare their respective fiscal years."""
    wb22 = openpyxl.load_workbook(F22_PATH, data_only=True)
    wb24 = openpyxl.load_workbook(F24_PATH, data_only=True)

    for sheet in wb22.sheetnames:
        t22 = wb22[sheet].cell(1, 2).value or wb22[sheet].cell(1, 1).value
        t24 = wb24[sheet].cell(1, 2).value or wb24[sheet].cell(1, 1).value

        assert "2022 - 23" in str(t22), f"Expected '2022 - 23' in {sheet} title: {t22}"
        assert "2023 - 24" in str(t24), f"Expected '2023 - 24' in {sheet} title: {t24}"

    wb22.close()
    wb24.close()
