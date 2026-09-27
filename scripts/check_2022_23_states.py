import openpyxl

wb22 = openpyxl.load_workbook('data/Road_EwayBill_2022_23.xlsx', data_only=True)

print("=== 2022-23 TABLE I ORIGIN VS DESTINATION ===")
ws1 = wb22['Tab I_Stat_to_Stat_Revised_Road']
origins = [ws1.cell(3, c).value for c in range(3, 36)]
dests = [ws1.cell(r, 2).value for r in range(4, 37)]

print(f"{'Index':<6} | {'Origin State (Col Header)':<25} | {'Dest State (Row Header)':<25} | {'Match?'}")
print("-" * 75)
for i, (orig, dest) in enumerate(zip(origins, dests), start=1):
    match = (orig == dest)
    flag = "MATCH" if match else "DIFF!"
    print(f"{i:<6} | {repr(orig):<25} | {repr(dest):<25} | {flag}")

print("\n=== 2022-23 STATE NAMES ACROSS ALL SHEETS ===")
ws3 = wb22['Tab III_Outward_Revised_Road']
ws4 = wb22['Tab IV_Inward_Revised_Road']
ws5 = wb22['Tab V_Internal_Revised_Road']

st3 = [ws3.cell(2, c).value for c in range(4, 37)]
st4 = [ws4.cell(2, c).value for c in range(4, 37)]
st5 = [ws5.cell(2, c).value for c in range(4, 37)]

print(f"{'Index':<6} | {'Tab I (Origin)':<22} | {'Tab I (Dest)':<22} | {'Tab III':<22} | {'Tab IV':<22} | {'Tab V':<22}")
print("-" * 115)
for i in range(33):
    print(f"{i+1:<6} | {str(origins[i]):<22} | {str(dests[i]):<22} | {str(st3[i]):<22} | {str(st4[i]):<22} | {str(st5[i]):<22}")
