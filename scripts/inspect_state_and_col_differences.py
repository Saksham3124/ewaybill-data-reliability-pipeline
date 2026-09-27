import openpyxl

wb22 = openpyxl.load_workbook('data/Road_EwayBill_2022_23.xlsx', data_only=True)
wb24 = openpyxl.load_workbook('data/Road_EwayBill_2023_24.xlsx', data_only=True)

print("=== CHECK TAB III COLUMNS IN 2022-23 ===")
ws3_22 = wb22['Tab III_Outward_Revised_Road']
ws3_24 = wb24['Tab III_Outward_Revised_Road']

print("2022-23 Tab III Col 37 (AK) header row 2:", repr(ws3_22.cell(2, 37).value))
print("2022-23 Tab III Col 37 (AK) row 3:", repr(ws3_22.cell(3, 37).value))
print("2022-23 Tab III Col 37 (AK) row 93:", repr(ws3_22.cell(93, 37).value))
print("2023-24 Tab III Col 37 (AK) header row 2:", repr(ws3_24.cell(2, 37).value))

print("\n=== STATE DIFFERENCES IN TAB I ===")
ws1_22 = wb22['Tab I_Stat_to_Stat_Revised_Road']
ws1_24 = wb24['Tab I_Stat_to_Stat_Revised_Road']

origin_22 = [ws1_22.cell(3, c).value for c in range(3, 36)]
dest_22 = [ws1_22.cell(r, 2).value for r in range(4, 37)]

origin_24 = [ws1_24.cell(3, c).value for c in range(3, 36)]
dest_24 = [ws1_24.cell(r, 2).value for r in range(4, 37)]

print("Origin 22 == Origin 24:", origin_22 == origin_24)
print("Dest 22 == Dest 24:", dest_22 == dest_24)

print("\nDifferences in Dest states (Row by Row):")
for r_idx, (d22, d24) in enumerate(zip(dest_22, dest_24), start=4):
    if d22 != d24:
        print(f"  Row {r_idx}: 2022-23={repr(d22)} vs 2023-24={repr(d24)}")

print("\nDest 22 set == Dest 24 set?:", set(dest_22) == set(dest_24))
if set(dest_22) != set(dest_24):
    print("In 22 but not 24:", set(dest_22) - set(dest_24))
    print("In 24 but not 22:", set(dest_24) - set(dest_22))

print("\n=== STATE DIFFERENCES IN TAB III, IV, V ===")
for sheet in ['Tab III_Outward_Revised_Road', 'Tab IV_Inward_Revised_Road', 'Tab V_Internal_Revised_Road']:
    ws22 = wb22[sheet]
    ws24 = wb24[sheet]
    states_22 = [ws22.cell(2, c).value for c in range(4, 37)]
    states_24 = [ws24.cell(2, c).value for c in range(4, 37)]
    print(f"\n[{sheet}]")
    print("States count: 22 =", len(states_22), ", 24 =", len(states_24))
    print("Sets equal?:", set(states_22) == set(states_24))
    print("List equal (ordering)?:", states_22 == states_24)
    if states_22 != states_24:
        for idx, (s22, s24) in enumerate(zip(states_22, states_24), start=4):
            if s22 != s24:
                col_letter = openpyxl.utils.get_column_letter(idx)
                print(f"  Col {col_letter} ({idx}): 2022-23={repr(s22)} vs 2023-24={repr(s24)}")

print("\n=== TITLES & HEADERS ACROSS ALL SHEETS ===")
for sheet in wb22.sheetnames:
    ws22 = wb22[sheet]
    ws24 = wb24[sheet]
    print(f"\n[{sheet}]")
    print("  2022-23 Row 1:", [ws22.cell(1, c).value for c in range(1, 10) if ws22.cell(1, c).value is not None])
    print("  2023-24 Row 1:", [ws24.cell(1, c).value for c in range(1, 10) if ws24.cell(1, c).value is not None])
    print("  2022-23 Row 2 (first 5 cols):", [ws22.cell(2, c).value for c in range(1, 6)])
    print("  2023-24 Row 2 (first 5 cols):", [ws24.cell(2, c).value for c in range(1, 6)])
